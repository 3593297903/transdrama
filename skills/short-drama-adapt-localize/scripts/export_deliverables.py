"""Export a single canonical text to TXT and optional DOCX; re-read both."""
import argparse
import json
from pathlib import Path
from drama_core import ContractError, load_json, save_new, validate


def canonical_lines(doc, review_draft=False):
    p = doc["project"]
    lines = [p["title"] + ("（待核审稿版）" if review_draft else ""),
             f"地区：{p['region']}；年代：{p['era']}；英语：{p['locale']}",
             "所有时长为制作规划；未进行音频实测或视频生成。9:16、2160×3840为输出目标。"]
    entities = {x["id"]: x for x in doc["entities"]}
    for seg in doc["segments"]:
        lines += [f"第{seg['episode']:02d}集｜{seg['id']}｜{seg['duration']}秒｜场次{seg['scene_id']}",
                  "原片时码：" + seg.get("source_timecode", "无（未提供有依据的原片时码）"), "①基础设定"]
        for ref in seg["asset_refs"]:
            asset = entities[ref["id"]]
            variant = next((x['description'] for x in asset.get('variants', []) if x['id'] == ref.get('variant_id')), "")
            description = asset['description'].rstrip("。； ")
            variant = variant.rstrip("。； ")
            lines.append(f"{asset['name']}：{description}" + (f"；本造型：{variant}" if variant else "") + f"；当前状态：{ref['state']}")
        lines += ["场景与空间：" + seg["setting"], "声音：" + seg["sound"], "②氛围与画质",
                  "9:16竖屏；2160×3840输出目标；" + seg["atmosphere"], "③分镜内容"]
        for shot in seg["shots"]:
            continuity = (f"；连续长镜{shot['continuous_take_id']}" if shot.get("continuous_take_id") else "")
            if shot.get("continues_shot_id"):
                continuity += f"；接续{shot['continues_shot_id']}，不是新切镜"
            lines.append(f"{shot['start']}—{shot['end']}秒｜{shot['id']}｜{shot['framing']}｜机位：{shot['camera']}｜构图：{shot['composition']}｜运镜：{shot['movement']}{continuity}。动作：{shot['action']}")
        # Audio is a single segment-level timeline, not repeated in every covered shot.
        for clip in seg["audio"]:
            label = "旁白" if clip["channel"] == "narration" else entities[clip["speaker_id"]]["name"] + ("（画外对白）" if clip["channel"] == "offscreen" else "")
            lines.append(f"声音 {clip['start']}—{clip['end']}秒｜{label}：{clip['text']}")
        lines.append("结束状态：" + seg["end_state"])
    lines += ["继承卡（不作为第二遍配音）"]
    for entity in doc["entities"]:
        lines.append(f"{entity.get('original_name', '原名未设')} → {entity['name']}（{entity['id']}，v{entity['version']}）")
        if entity.get("relationship"):
            lines.append("关系：" + entity["relationship"])
    if doc["segments"]:
        lines.append("最新结尾：" + doc["segments"][-1]["end_state"])
    lines.append("待核问题：" + ("见同目录 validation.json；本版仅供审稿。" if review_draft else "无未关闭的已记录文字问题；英语为单模型审读，非母语盲评。"))
    return lines


def _create_docx(lines, path):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.8)
    section.left_margin = section.right_margin = Inches(.85)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    fonts = normal.element.get_or_add_rPr()
    east = OxmlElement("w:rFonts")
    east.set(qn("w:eastAsia"), "Microsoft YaHei")
    fonts.append(east)
    for i, line in enumerate(lines):
        para = document.add_paragraph(line)
        if i == 0:
            para.runs[0].bold = True
            para.runs[0].font.size = Pt(18)
        elif line.startswith("第") and "｜" in line:
            para.runs[0].bold = True
            para.paragraph_format.keep_with_next = True
        elif line in ("①基础设定", "②氛围与画质", "③分镜内容", "继承卡（不作为第二遍配音）"):
            para.runs[0].bold = True
            para.paragraph_format.keep_with_next = True
    document.save(path)
    actual = [p.text for p in Document(path).paragraphs]
    if actual != lines:
        raise ContractError("DOCX read-back differs from canonical text")


def export(doc, out, want_docx=False, allow_review_draft=False, docx_writer=None):
    report = validate(doc)
    if report["errors"]:
        raise ContractError("Structural validation failed; run check_deliverables.py for details")
    if report["review_required"] and not allow_review_draft:
        raise ContractError("Semantic/reading review pending; resolve or explicitly export review draft")
    out = Path(out)
    if out.exists():
        raise ContractError(f"Refusing existing output directory: {out}")
    out.mkdir(parents=True)
    lines = canonical_lines(doc, bool(report["review_required"]))
    txt = out / "完整视频提示词.txt"
    with txt.open("x", encoding="utf-8", newline="\n") as stream:
        stream.write("\n".join(lines) + "\n")
    if txt.read_text(encoding="utf-8") != "\n".join(lines) + "\n":
        raise ContractError("TXT read-back differs")
    status = {"txt": "created_and_read_back", "docx": "not_requested",
              "docx_visual_qa": "not_performed_by_exporter", "validation": report["overall_status"]}
    if want_docx:
        try:
            (docx_writer or _create_docx)(lines, out / "完整视频提示词.docx")
            status["docx"] = "created_and_read_back"
        except ImportError:
            status["docx"] = "unavailable_dependency_txt_delivered"
        except Exception as exc:
            status["docx"] = "failed"
            status["docx_error"] = str(exc)
            save_new(out / "export-status.json", status)
            save_new(out / "validation.json", report)
            raise ContractError("DOCX export failed; TXT retained, see export-status.json") from exc
    save_new(out / "validation.json", report)
    save_new(out / "export-status.json", status)
    return status


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("project")
    parser.add_argument("--out", required=True)
    parser.add_argument("--docx", action="store_true")
    parser.add_argument("--allow-review-draft", action="store_true")
    args = parser.parse_args()
    try:
        status = export(load_json(args.project), args.out, args.docx, args.allow_review_draft)
        print(json.dumps(status, ensure_ascii=False, indent=2))
        return 3 if status["validation"] == "review_draft" or status["docx"].startswith("unavailable") else 0
    except (OSError, ValueError) as exc:
        print(f"ERROR: {exc}")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
